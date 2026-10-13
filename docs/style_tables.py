#!/usr/bin/env python3
"""Give pandoc-generated tables the look of the stationary test report:
light grey grid, blue-shaded bold header row, full text width, 9.5 pt text.

Usage: python docs/style_tables.py FILE.docx [FILE.docx ...]   (edits in place)
"""
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

GRID, HEADER_FILL = 'BFBFBF', 'DCE6F1'


def set_child(parent, tag, **attrs):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for k, v in attrs.items():
        el.set(qn(k), v)
    return el


def style_table(table):
    tblPr = table._tbl.tblPr
    set_child(tblPr, 'w:tblW', **{'w:w': '5000', 'w:type': 'pct'})
    borders = tblPr.find(qn('w:tblBorders'))
    if borders is not None:
        tblPr.remove(borders)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{side}')
        for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '0'), ('w:color', GRID)):
            b.set(qn(k), v)
        borders.append(b)
    tblPr.append(borders)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            if i == 0:
                set_child(cell._tc.get_or_add_tcPr(), 'w:shd',
                          **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': HEADER_FILL})
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    if i == 0:
                        r.bold = True


for path in sys.argv[1:]:
    doc = docx.Document(path)
    for t in doc.tables:
        style_table(t)
    doc.save(path)
